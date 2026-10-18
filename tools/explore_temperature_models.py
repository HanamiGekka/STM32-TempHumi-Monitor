"""Compare static calibration families without treating transient pairs as truth.

All tuning uses entire held-out experiment phases; no random adjacent-row split.
The original workbook is read only. Predictions outside a training range are
reported separately from in-range results and are never deployment approval.
"""
from pathlib import Path
import argparse, csv, hashlib, json, warnings
import numpy as np
from scipy import interpolate, optimize
from sklearn.isotonic import IsotonicRegression
from analyze_temperature_calibration import metrics, PHASES, weight
from openpyxl import load_workbook


def load(path):
    book=load_workbook(path,read_only=True,data_only=False);sheet=book['Sheet1'];records=[]
    for row in range(2,51):
        x,y,note=(sheet.cell(row,c).value for c in (1,2,3))
        if not all(isinstance(v,(int,float)) and not isinstance(v,bool) and np.isfinite(v) for v in (x,y)):
            raise ValueError('Missing/nonfinite numeric pair at Excel row '+str(row))
        phase=next(label for start,end,label in PHASES if start<=row<=end)
        records.append({'excel_row':row,'dut_deg_c':float(x),'reference_deg_c':float(y),'phase':phase,'original_note':note or ''})
    book.close();return records


def basis(x, knots):
    """Linear hat basis; linear extrapolation is diagnostic only."""
    x=np.asarray(x);k=np.asarray(knots);a=np.zeros((len(x),len(k)))
    i=np.clip(np.searchsorted(k,x)-1,0,len(k)-2)
    f=(x-k[i])/(k[i+1]-k[i]);a[np.arange(len(x)),i]=1-f;a[np.arange(len(x)),i+1]=f
    return a


def model_specs():
    s=[('baseline',{}),('offset_mean',{}),('offset_median',{}),('affine_l2',{}),
       ('affine_l1',{}),('affine_huber',{}),('affine_soft_l1',{})]
    s += [('polynomial',{'degree':d}) for d in (2,3,4,5)]
    s += [('piecewise',{'segments':n,'loss':loss}) for n in (2,3,4,6) for loss in ('linear','huber','soft_l1','l1')]
    s += [('monotone_knots',{'segments':n,'loss':loss,'smooth':lam})
          for n in (2,3,4,6) for loss in ('linear','huber','l1') for lam in (0.,1.)]
    s += [('isotonic',{}),('isotonic_pchip',{})]
    s += [('smoothing_spline',{'s':v}) for v in (1.,10.,30.,60.)]
    s += [('local_linear',{'fraction':v}) for v in (.35,.6,.85)]
    s += [('segmented_offset',{'segments':v}) for v in (2,3,4)]
    s += [('anchored_monotone',{'segments':n,'loss':loss}) for n in (3,4,6) for loss in ('linear','huber','l1')]
    s += [('guarded_piecewise',{'segments':n,'loss':loss}) for n in (2,3,4,6) for loss in ('linear','huber','l1')]
    return [{'family':f,**p,'id':f+(':'+','.join(str(v) for v in p.values()) if p else '')} for f,p in s]


def fit(spec,x,y):
    f=spec['family'];x=np.asarray(x);y=np.asarray(y);lo=float(min(x));hi=float(max(x))
    m={'spec':spec,'training_domain':[lo,hi]}
    if f=='baseline':m['kind']='offset';m['offset']=0.
    elif f.startswith('offset_'):
        m.update(kind='offset',offset=float(np.mean(y-x) if f.endswith('mean') else np.median(y-x)))
    elif f.startswith('affine') or f=='polynomial':
        d=spec.get('degree',1);center=float(np.mean(x));scale=max(float(np.std(x)),.1)
        a=np.polynomial.polynomial.polyvander((x-center)/scale,d)
        if f=='affine_l1':
            n=len(x);o=np.r_[np.zeros(d+1),np.ones(n)]
            res=optimize.linprog(o,A_ub=np.r_[np.c_[a,-np.eye(n)],np.c_[-a,-np.eye(n)]],b_ub=np.r_[y,-y],bounds=[(None,None)]*(d+1)+[(0,None)]*n,method='highs')
            if not res.success:raise ValueError(res.message)
            c=res.x[:d+1]
        else:
            c=np.linalg.lstsq(a,y,rcond=None)[0]
            if f in ('affine_huber','affine_soft_l1'):
                c=optimize.least_squares(lambda z:a@z-y,c,loss=f[7:],f_scale=1.).x
        m.update(kind='poly',center=center,scale=scale,coefficients=c.tolist())
    elif f=='guarded_piecewise':
        # Deployment support is fixed from observed input range, independent of
        # reference labels. No correction is applied outside this support.
        k=np.linspace(25.2,35.6,spec['segments']+1);a=weight(x)[:,None]*basis(x,k)
        grid=np.unique(np.r_[np.linspace(25.2,35.6,1041),k+1e-6,k-1e-6])
        delta=1e-7
        deriv=(weight(grid+delta)[:,None]*basis(grid+delta,k)-weight(grid-delta)[:,None]*basis(grid-delta,k))/(2*delta)
        n=len(k);target=y-x;loss=spec['loss']
        if loss=='l1':
            aub=np.r_[np.c_[a,-np.eye(len(x))],np.c_[-a,-np.eye(len(x))],np.c_[deriv,np.zeros((len(grid),len(x)))],np.c_[-deriv,np.zeros((len(grid),len(x)))]]
            bub=np.r_[target,-target,np.full(len(grid),3.),np.full(len(grid),.95)]
            sol=optimize.linprog(np.r_[np.zeros(n),np.ones(len(x))],A_ub=aub,b_ub=bub,bounds=[(0,6)]*n+[(0,None)]*len(x),method='highs')
        else:
            def obj(c):
                e=a@c-target
                return np.sum(e*e) if loss=='linear' else np.sum(np.where(abs(e)<=1.,.5*e*e,abs(e)-.5))
            sol=optimize.minimize(obj,np.full(n,.5),bounds=[(0,6)]*n,constraints=[optimize.LinearConstraint(deriv,-.95,3.)],method='SLSQP',options={'maxiter':1000,'ftol':1e-10})
        if not sol.success:raise ValueError(sol.message)
        m.update(kind='guarded',knots=k.tolist(),offsets=sol.x[:n].tolist(),domain=[25.2,35.6],slope_constraints=[.05,4.])
    elif f in ('piecewise','monotone_knots','anchored_monotone'):
        n=spec['segments'];k=np.unique(np.quantile(x,np.linspace(0,1,n+1)))
        a=basis(x,k);n=len(k);loss=spec['loss'];lam=spec.get('smooth',0.)
        diff=np.diff(np.eye(n),n=2,axis=0)
        constrained=f!='piecewise'
        initial=np.interp(k,np.sort(x),y[np.argsort(x)])
        if constrained:initial=k.copy()+min(float(np.median(y-x)),1.)
        slopes=np.diff(np.eye(n),axis=0)
        # Positive knot slopes prevent inversions; endpoint identity is only for
        # the explicitly deployment-guarded family. It is not fake training data.
        cons=[]
        if constrained:cons=[optimize.LinearConstraint(slopes,.05*np.diff(k),4.*np.diff(k))]
        bounds=[(None,None)]*n
        if f=='anchored_monotone':bounds[0]=(lo,lo);bounds[-1]=(hi,hi);initial=k.copy()
        if loss=='l1' and lam==0:
            ob=np.r_[np.zeros(n),np.ones(len(x))];aub=np.r_[np.c_[a,-np.eye(len(x))],np.c_[-a,-np.eye(len(x))]];bub=np.r_[y,-y]
            if constrained:
                aub=np.r_[aub,np.c_[-slopes,np.zeros((n-1,len(x)))],np.c_[slopes,np.zeros((n-1,len(x)))]]
                bub=np.r_[bub,-.05*np.diff(k),4.*np.diff(k)]
            res=optimize.linprog(ob,A_ub=aub,b_ub=bub,bounds=bounds+[(0,None)]*len(x),method='highs')
            if not res.success:raise ValueError(res.message)
            c=res.x[:n]
        else:
            def objective(c):
                e=a@c-y
                cost=np.sum(e*e) if loss=='linear' else (np.sum(abs(e)) if loss=='l1' else np.sum(np.where(abs(e)<=1.,.5*e*e,abs(e)-.5)))
                return cost+lam*np.sum((diff@c)**2)
            if constrained or loss=='l1':
                res=optimize.minimize(objective,initial,method='SLSQP',constraints=cons,bounds=bounds,options={'maxiter':600,'ftol':1e-9})
                if not res.success:raise ValueError(res.message)
                c=res.x
            else:
                c=optimize.least_squares(lambda z:a@z-y,np.linalg.lstsq(a,y,rcond=None)[0],loss=loss,f_scale=1.).x
        m.update(kind='knots',knots=k.tolist(),values=c.tolist(),anchored=f=='anchored_monotone')
    elif f.startswith('isotonic'):
        iso=IsotonicRegression(increasing=True,out_of_bounds='clip').fit(x,y)
        m.update(kind='pchip' if f.endswith('pchip') else 'knots',knots=iso.X_thresholds_.tolist(),values=iso.y_thresholds_.tolist())
    elif f=='smoothing_spline':
        ux=np.unique(x);uy=np.array([np.mean(y[x==v]) for v in ux]);counts=np.array([sum(x==v) for v in ux])
        spline=interpolate.UnivariateSpline(ux,uy,w=np.sqrt(counts),s=spec['s'],k=min(3,len(ux)-1))
        m.update(kind='spline',tck=[z.tolist() if hasattr(z,'tolist') else int(z) for z in spline._eval_args])
    elif f=='segmented_offset':
        cuts=np.unique(np.quantile(x,np.linspace(0,1,spec['segments']+1)))[1:-1];labels=np.searchsorted(cuts,x)
        offsets=[float(np.median((y-x)[labels==i])) if np.any(labels==i) else float(np.median(y-x)) for i in range(len(cuts)+1)]
        m.update(kind='step_offset',cuts=cuts.tolist(),offsets=offsets)
    elif f=='local_linear':m.update(kind='local',x=x.tolist(),y=y.tolist(),fraction=spec['fraction'])
    else:raise ValueError(f)
    return m


def predict(m,z):
    z=np.asarray(z,dtype=float);kind=m['kind']
    if kind=='offset':return z+m['offset']
    if kind=='guarded':return z+weight(z)*(basis(z,m['knots'])@m['offsets'])
    if kind=='poly':return np.polynomial.polynomial.polyval((z-m['center'])/m['scale'],m['coefficients'])
    if kind=='knots':
        p=basis(z,m['knots'])@m['values']
        if m.get('anchored'):p=np.where((z<m['training_domain'][0])|(z>m['training_domain'][1]),z,p)
        return p
    if kind=='pchip':return interpolate.PchipInterpolator(m['knots'],m['values'],extrapolate=True)(z)
    if kind=='spline':return interpolate.splev(z,m['tck'])
    if kind=='step_offset':return z+np.asarray(m['offsets'])[np.searchsorted(m['cuts'],z)]
    if kind=='local':
        x=np.asarray(m['x']);y=np.asarray(m['y']);out=[];k=max(3,int(np.ceil(len(x)*m['fraction'])))
        for point in z:
            distance=abs(x-point);width=max(float(np.sort(distance)[min(k-1,len(x)-1)]),.01)*1.00001
            w=np.maximum(0,1-(distance/width)**3)**3;w[distance>width]=0
            a=np.c_[np.ones(len(x)),x-point];c=np.linalg.lstsq(a*np.sqrt(w[:,None]),y*np.sqrt(w),rcond=None)[0];out.append(c[0])
        return np.asarray(out)
    raise ValueError(kind)


def evaluate(spec,x,y,rows):
    m=fit(spec,x,y);p=predict(m,x);grid=np.linspace(min(x),max(x),2081);curve=predict(m,grid);step=.00001
    phase_results={}
    for start,end,label in PHASES:
        mask=(rows>=start)&(rows<=end);phase_results[label]=metrics(p[mask],y[mask])
    cv=[]
    for start,end,label in PHASES:
        test=(rows>=start)&(rows<=end);train=~test
        try:
            fm=fit(spec,x[train],y[train]);fp=predict(fm,x[test]);inside=(x[test]>=min(x[train]))&(x[test]<=max(x[train]))
            cv.append({'phase':label,'training_rows':rows[train].tolist(),'held_out_rows':rows[test].tolist(),'metrics':metrics(fp,y[test]),'predictions':fp.tolist(),'in_training_domain':inside.tolist(),'in_range_metrics':metrics(fp[inside],y[test][inside]) if any(inside) else None})
        except Exception as exc:cv.append({'phase':label,'error':str(exc)})
    if any('error' in v for v in cv):pooled=None
    else:pooled=metrics(np.concatenate([v['predictions'] for v in cv]),np.concatenate([y[(rows>=a)&(rows<=b)] for a,b,_ in PHASES]))
    chronology=[]
    for end in (8,19,28,36):
        train=rows<=end;test=~train
        try:
            fm=fit(spec,x[train],y[train]);fp=predict(fm,x[test]);inside=(x[test]>=min(x[train]))&(x[test]<=max(x[train]))
            chronology.append({'fit_through_row':end,'metrics':metrics(fp,y[test]),'in_range_metrics':metrics(fp[inside],y[test][inside]) if any(inside) else None,'in_range_count':int(sum(inside))})
        except Exception as exc:chronology.append({'fit_through_row':end,'error':str(exc)})
    d=np.diff(curve)/np.diff(grid)
    left=predict(m,np.array([min(x)-step,min(x)+step]));right=predict(m,np.array([max(x)-step,max(x)+step]))
    return {'id':spec['id'],'spec':spec,'model':m,'full':metrics(p,y),'phase_metrics':phase_results,'leave_phase_out':cv,'pooled_phase_holdout':pooled,'chronological':chronology,
            'shape':{'minimum_slope_grid':float(min(d)),'maximum_slope_grid':float(max(d)),'monotonic_grid':bool(min(d)>=-1e-7),'boundary_change_over_2e_minus5':[float(left[1]-left[0]),float(right[1]-right[0])],'at_0_20_25_35_50':predict(m,np.array([0,20,25,35,50])).tolist(),'note':'OUT-OF-DOMAIN VALUES ARE DIAGNOSTIC, NOT DEPLOYABLE CALIBRATION'},'predictions':p.tolist()}


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--workbook',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=True)
    before=hashlib.sha256(args.workbook.read_bytes()).hexdigest();records=load(args.workbook);rows=np.array([v['excel_row'] for v in records]);x=np.array([v['dut_deg_c'] for v in records]);y=np.array([v['reference_deg_c'] for v in records])
    results=[];failed=[]
    for i,s in enumerate(model_specs()):
        try:
            with warnings.catch_warnings():
                warnings.simplefilter('error',RuntimeWarning);result=evaluate(s,x,y,rows)
            results.append(result);print(i+1,s['id'],'full',round(result['full']['mae_deg_c'],4),'LOPO',round(result['pooled_phase_holdout']['mae_deg_c'],4) if result['pooled_phase_holdout'] else 'FAILED',flush=True)
        except Exception as exc:failed.append({'spec':s,'error':str(exc)});print('FAILED',s['id'],str(exc),flush=True)
    deployed=x+.9*weight(x)
    output={'workbook_sha256':before,'classification':'INTERNAL EXPLORATION / phase CV; selected-family scores are not unbiased independent accuracy estimates','original_rows':records,'models':results,'failed_models':failed,'deployed_provisional':{'full':metrics(deployed,y),'predictions':deployed.tolist()},'excluded_rows':[],'physical_parameter_fit':'Not identifiable from paired temperatures without verified actual R, excitation/gain and synchronized equilibrium reference. Datasheet bounds support sensitivity analysis only.'}
    (args.output/'model_sweep.json').write_text(json.dumps(output,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    with (args.output/'model_ranking.csv').open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.DictWriter(f,fieldnames=['id','full_mae','full_rmse','full_max','ambient_mae','phase_cv_mae','phase_cv_max','monotonic','min_slope','max_slope']);writer.writeheader()
        for v in sorted(results,key=lambda v:v['full']['mae_deg_c']):
            cv=v['pooled_phase_holdout'];writer.writerow({'id':v['id'],'full_mae':v['full']['mae_deg_c'],'full_rmse':v['full']['rmse_deg_c'],'full_max':v['full']['max_abs_deg_c'],'ambient_mae':v['phase_metrics']['Ambient']['mae_deg_c'],'phase_cv_mae':cv['mae_deg_c'] if cv else None,'phase_cv_max':cv['max_abs_deg_c'] if cv else None,'monotonic':v['shape']['monotonic_grid'],'min_slope':v['shape']['minimum_slope_grid'],'max_slope':v['shape']['maximum_slope_grid']})
    assert hashlib.sha256(args.workbook.read_bytes()).hexdigest()==before
    print('COMPLETE',len(results),'models;',len(failed),'failed',flush=True)


if __name__=='__main__':main()
