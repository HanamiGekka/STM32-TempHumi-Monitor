"""Reproduce RUN 3.1 internal comparison analysis without rewriting its workbook."""
from pathlib import Path
import argparse, collections, csv, hashlib, json, math, statistics
import numpy as np
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
PHASES = [(2,8,'Ambient'),(9,19,'Breath'),(20,28,'Warm water'),
          (29,36,'Hotter water'),(37,50,'Further heating')]
LOW, FULL_LOW, FULL_HIGH, HIGH = 25.2, 25.7, 32.6, 35.6

def weight(t):
    """C1 smooth envelope inside observed DUT range; no correction outside it."""
    t = np.asarray(t,dtype=float)
    lo = np.clip((t-LOW)/(FULL_LOW-LOW),0,1)
    hi = np.clip((HIGH-t)/(HIGH-FULL_HIGH),0,1)
    return lo*lo*(3-2*lo)*hi*hi*(3-2*hi)

def metrics(dut, reference):
    e = np.asarray(dut)-np.asarray(reference)
    return {'n':len(e),'bias_deg_c':float(np.mean(e)),
            'mae_deg_c':float(np.mean(abs(e))), 'rmse_deg_c':float(np.sqrt(np.mean(e*e))),
            'max_abs_deg_c':float(max(abs(e))),
            **{f'within_{limit:g}_deg_c':{'count':int(np.sum(abs(e)<=limit+1e-9)),
              'pct':float(np.mean(abs(e)<=limit+1e-9)*100)} for limit in (0.5,1,2)}}

def fit(name,x,y):
    if name=='A_baseline':return (1.,0.)
    if name in ('B_robust_offset','B_bounded_offset'):return (1.,float(np.median(y-x)))
    if name=='B_mean_offset':return (1.,float(np.mean(y-x)))
    return tuple(float(v) for v in np.polyfit(x,y,1))

def predict(name,coeff,x):
    a,b = coeff
    return x+b*weight(x) if name=='B_bounded_offset' else a*x+b

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--workbook',type=Path,default=ROOT/'references/温度采集数据对照表.xlsx')
    p.add_argument('--output',type=Path,required=True);args=p.parse_args()
    out=args.output;out.mkdir(parents=True,exist_ok=True)
    before=hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    workbook=load_workbook(args.workbook,read_only=True,data_only=False)
    sheet=workbook['Sheet1'];records=[]
    for row in range(2,sheet.max_row+1):
        x,y,note=(sheet.cell(row,c).value for c in (1,2,3))
        if not all(isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v) for v in (x,y)):
            raise ValueError('Non-numeric/missing pair at Excel row '+str(row))
        phase=next(label for start,end,label in PHASES if start<=row<=end)
        records.append({'excel_row':row,'dut_deg_c':float(x),'reference_deg_c':float(y),
                        'phase':phase,'original_note':note or ''})
    workbook.close();assert len(records)==49
    x=np.array([r['dut_deg_c'] for r in records]);y=np.array([r['reference_deg_c'] for r in records]);rows=np.array([r['excel_row'] for r in records])
    baseline=metrics(x,y)
    assert (min(x),max(x),min(y),max(y))==(25.2,35.6,25.9,44.5)
    assert abs(baseline['bias_deg_c']+1.8632653061224487)<1e-10
    assert abs(baseline['rmse_deg_c']-2.7187557177983503)<1e-10
    assert baseline['within_1_deg_c']['count']==18
    models=['A_baseline','B_robust_offset','B_mean_offset','C_affine','B_bounded_offset']
    exploratory={}
    for name in models:
        c=fit(name,x,y);exploratory[name]={'coefficients':list(c),'metrics':metrics(predict(name,c,x),y)}
    validation=[]
    splits=[('ambient chronological',rows<=6,(rows>=7)&(rows<=8)),
            ('ambient to heating',rows<=8,rows>=9),('earlier to late heating',rows<=36,rows>=37)]
    splits += [(f'leave phase out: {label}',~((rows>=start)&(rows<=end)),(rows>=start)&(rows<=end)) for start,end,label in PHASES]
    for label,train,test in splits:
        for name in models:
            c=fit(name,x[train],y[train]);validation.append({'split':label,'model':name,
                'training_excel_rows':rows[train].tolist(),'validation_excel_rows':rows[test].tolist(),
                'coefficients':list(c),'training':metrics(predict(name,c,x[train]),y[train]),
                'held_out':metrics(predict(name,c,x[test]),y[test])})
    # Fit only initial ambient rows 2..6. Rows 7..8 remain an internal ambient check.
    train=rows<=6; b=float(np.round(np.median(y[train]-x[train]),6));assert b==0.9
    corrected=x+b*weight(x)
    selected={'model':'B_bounded_offset','offset_deg_c':b,
              'training_excel_rows':rows[train].tolist(),'held_out_excel_rows':rows[~train].tolist(),
              'domain_policy':{'no_correction_at_or_outside':[LOW,HIGH],
                'full_offset_range':[FULL_LOW,FULL_HIGH],'transition':'cubic smoothstep inside evidence domain',
                'max_absolute_offset_deg_c':b,'minimum_derivative':1-b*1.5/(HIGH-FULL_HIGH)},
              'training':metrics(corrected[train],y[train]),'held_out':metrics(corrected[~train],y[~train]),
              'full':metrics(corrected,y),'phases':{label:{'raw':metrics(x[(rows>=s)&(rows<=e)],y[(rows>=s)&(rows<=e)]),
                 'corrected':metrics(corrected[(rows>=s)&(rows<=e)],y[(rows>=s)&(rows<=e)])} for s,e,label in PHASES}}
    for i,record in enumerate(records):
        record.update(raw_error_deg_c=float(x[i]-y[i]),correction_deg_c=float(corrected[i]-x[i]),
          corrected_deg_c=float(corrected[i]),corrected_error_deg_c=float(corrected[i]-y[i]),
          partition='ambient fit' if train[i] else 'internal holdout')
    with (out/'comparison_rows.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(records[0]));w.writeheader();w.writerows(records)
    duplicates=[]
    for value in sorted(set(x)):
        ix=np.flatnonzero(x==value)
        if len(ix)>1 and max(y[ix])-min(y[ix])>0:
            duplicates.append({'dut_deg_c':float(value),'excel_rows':rows[ix].tolist(),
                               'reference_deg_c':y[ix].tolist(),'irreducible_static_max_error_lower_bound_deg_c':float((max(y[ix])-min(y[ix]))/2)})
    influence=[]
    for index in np.argsort(abs(x-y))[-5:][::-1]:
        keep=np.arange(len(x))!=index
        influence.append({'excluded_for_sensitivity_only_excel_row':int(rows[index]),
                          'affine_without_row':list(fit('C_affine',x[keep],y[keep])),
                          'raw_without_row':metrics(x[keep],y[keep])})
    grid=np.linspace(-40,125,33001);g=grid+b*weight(grid)
    check={'monotonic':bool(np.all(np.diff(g)>0)), 'max_offset_deg_c':float(np.max(abs(g-grid))),
           'outside_domain_exact_baseline':bool(np.all(g[(grid<=LOW)|(grid>=HIGH)]==grid[(grid<=LOW)|(grid>=HIGH)])),
           'ambient_holdout_improves':metrics(corrected[(rows>=7)&(rows<=8)],y[(rows>=7)&(rows<=8)])['mae_deg_c']<metrics(x[(rows>=7)&(rows<=8)],y[(rows>=7)&(rows<=8)])['mae_deg_c'],
           'all_heating_phases_mae_improve':all(v['corrected']['mae_deg_c']<v['raw']['mae_deg_c'] for k,v in selected['phases'].items() if k!='Ambient')}
    result={'workbook_sha256':before,'worksheet':'Sheet1','rows':49,'dut_range_deg_c':[float(min(x)),float(max(x))],
            'reference_range_deg_c':[float(min(y)),float(max(y))], 'baseline':baseline,'exploratory':exploratory,
            'validation_classification':'INTERNAL VALIDATION; phase/chronology based; no independent acquisition',
            'validation':validation,'selected':selected,'checks':check,
            'model_D':'NOT IDENTIFIABLE: no verified as-built circuit/NTC values or paired constituent voltage/raw data in workbook; no physical constant fit performed',
            'duplicates':duplicates,'high_error_influence':influence,
            'reference_status':'DELIXI THM-01 working reference; user-supplied +/-0.3 deg C is not a calibration certificate/uncertainty',
            'accuracy':'NOT INDEPENDENTLY VERIFIED; +/-1 deg C over 0..50 deg C NOT ESTABLISHED',
            'excluded_observations':[]}
    (out/'analysis.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False})
    fig,axes=plt.subplots(2,3,figsize=(15,8),layout='constrained')
    for j,(start,end,label) in enumerate(PHASES):
        ix=(rows>=start)&(rows<=end);color=f'C{j}'
        axes[0,0].scatter(x[ix],y[ix],label=label,color=color,s=30)
        axes[0,1].scatter(x[ix],x[ix]-y[ix],color=color,s=30)
        axes[0,2].scatter(y[ix],corrected[ix]-y[ix],color=color,s=30)
    axes[0,0].plot([24,46],[24,46],color='gray',linestyle='--');axes[0,0].legend(fontsize=8)
    axes[0,0].set(xlabel='DUT (deg C)',ylabel='DELIXI working reference (deg C)',title='49 manual pairs\nPhases are not equilibrium proof')
    axes[0,1].set(xlabel='DUT (deg C)',ylabel='DUT - reference (deg C)',title='Raw residual versus DUT')
    axes[0,2].set(xlabel='Reference (deg C)',ylabel='Corrected - reference (deg C)',title='Provisional residual versus reference')
    axes[1,0].plot(rows,x,label='DUT');axes[1,0].plot(rows,y,label='Reference');axes[1,0].plot(rows,corrected,label='Provisional correction');axes[1,0].legend(fontsize=8)
    axes[1,0].set(xlabel='Excel row (NOT elapsed time)',ylabel='deg C',title='Sequence/thermal mismatch; timestamps absent')
    axes[1,1].plot(rows,x-y,label='Raw');axes[1,1].plot(rows,corrected-y,label='Corrected');axes[1,1].legend(fontsize=8)
    axes[1,1].set(xlabel='Excel row',ylabel='Error (deg C)',title='All observations retained, including row 50')
    domain=np.linspace(20,50,601);axes[1,2].plot(domain,b*weight(domain));axes[1,2].axvspan(LOW,HIGH,alpha=.1,color='gray')
    axes[1,2].set(xlabel='Uncorrected temperature (deg C)',ylabel='Applied offset (deg C)',title='C1 monotonic transfer\nZero correction outside observed domain')
    for ax in [axes[0,1],axes[0,2],axes[1,1]]:
        ax.axhline(0,color='gray',lw=.8);ax.axhline(1,color='gray',lw=.6,ls=':');ax.axhline(-1,color='gray',lw=.6,ls=':')
    fig.savefig(out/'comparison.png',dpi=160);plt.close(fig)
    assert hashlib.sha256(args.workbook.read_bytes()).hexdigest()==before
    assert all(check[k] for k in ('monotonic','outside_domain_exact_baseline','ambient_holdout_improves','all_heating_phases_mae_improve'))
    print(json.dumps({'baseline':baseline,'selected':selected,'checks':check},indent=2))

if __name__=='__main__':main()
